"""Banco de teses (19/10/2026, pedido do Defensor).

Tópicos de petição de cunho GENÉRICO, com teses jurídicas aplicáveis a casos
concretos, cadastrados na página "Banco de teses" (menu, depois de Agenda).
Quando uma petição é gerada por IA com o checkbox "Usar banco de teses"
marcado, as teses ATIVAS vão no prompt (`bloco_prompt`): a IA confere em cada
uma se o "quando se aplica" bate com o caso e, se bater, escreve só a marca
`[[TESE:id]]` (parágrafo sozinho) no ponto em que o tópico entra. O PAINEL troca
a marca pelo texto cadastrado, IPSIS LITTERIS (`inserir_teses_html` /
`inserir_teses_docx`), em AZUL — indicativo visual de que o trecho veio do
banco; o Defensor muda a cor no próprio editor ao revisar. A tese nunca passa
pela redação da IA (20/10/2026: na 1ª versão a IA "adaptava" e acabava
resumindo e cortando ementas/artigos). A IA NÃO lista as teses
aplicadas/descartadas na resposta (o azul já é o indicativo).

Armazenamento: `dados/banco_teses.json` (lista de teses).
"""

from __future__ import annotations

import contextlib
import json
import re
import secrets
import threading
from datetime import datetime
from pathlib import Path

_STORE = Path(__file__).resolve().parents[1].parent / "dados" / "banco_teses.json"
_lock = threading.Lock()

TESES_DPU_DIR = Path(r"C:\Claude\teses-dpu")

# Cor do trecho vindo do banco: o azul da paleta do editor (Quill), para o
# Defensor reconhecê-lo/trocá-lo pelo mesmo seletor de cor.
COR_TESE_HEX = "#0066cc"
COR_TESE_RGB = (0, 102, 204)

AREAS = [
    {"chave": "geral", "rotulo": "Geral (qualquer área)"},
    {"chave": "beneficio-incapacidade", "rotulo": "Benefício por incapacidade"},
    {"chave": "bpc-loas", "rotulo": "BPC/LOAS"},
    {"chave": "aposentadoria", "rotulo": "Aposentadoria"},
    {"chave": "pensao", "rotulo": "Pensão por morte"},
    {"chave": "auxilio-reclusao", "rotulo": "Auxílio-reclusão"},
    {"chave": "salario-maternidade", "rotulo": "Salário-maternidade"},
    {"chave": "penal-militar", "rotulo": "Penal militar (JMU)"},
    {"chave": "outra", "rotulo": "Outra"},
]
_AREAS_VALIDAS = {a["chave"] for a in AREAS}


# ---- armazenamento -------------------------------------------------------------

def _ler() -> list[dict]:
    with contextlib.suppress(Exception):
        if _STORE.exists():
            dados = json.loads(_STORE.read_text(encoding="utf-8"))
            if isinstance(dados, list):
                return dados
    return []


def _gravar(teses: list[dict]) -> None:
    _STORE.parent.mkdir(parents=True, exist_ok=True)
    tmp = _STORE.with_suffix(".tmp")
    tmp.write_text(json.dumps(teses, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp.replace(_STORE)


def listar() -> list[dict]:
    return sorted(_ler(), key=lambda t: (not t.get("ativa", True), (t.get("titulo") or "").lower()))


def obter(tese_id: str) -> dict | None:
    return next((t for t in _ler() if t.get("id") == tese_id), None)


def _limpar(dados: dict) -> dict:
    areas = [a for a in (dados.get("areas") or []) if a in _AREAS_VALIDAS] or ["geral"]
    return {
        "titulo": (dados.get("titulo") or "").strip()[:200],
        "areas": areas,
        "quando_aplica": (dados.get("quando_aplica") or "").strip()[:4000],
        "texto_html": (dados.get("texto_html") or "").strip()[:60000],
        "fundamentos": (dados.get("fundamentos") or "").strip()[:4000],
        "ativa": bool(dados.get("ativa", True)),
    }


def salvar(dados: dict) -> dict:
    """Cria (sem `id`) ou atualiza (com `id`) uma tese. Devolve a tese gravada."""
    limpo = _limpar(dados or {})
    if not limpo["titulo"]:
        raise ValueError("Informe o título da tese.")
    if not re.sub(r"<[^>]+>", "", limpo["texto_html"]).strip():
        raise ValueError("Informe o texto do tópico.")
    agora = datetime.now().isoformat(timespec="seconds")
    with _lock:
        teses = _ler()
        tese_id = (dados or {}).get("id") or ""
        atual = next((t for t in teses if t.get("id") == tese_id), None) if tese_id else None
        if atual:
            atual.update(limpo, atualizado_em=agora)
            tese = atual
        else:
            tese = {"id": "t" + secrets.token_hex(4), **limpo, "criado_em": agora,
                    "atualizado_em": agora, "origem": (dados or {}).get("origem") or "manual"}
            teses.append(tese)
        _gravar(teses)
    return tese


def excluir(tese_id: str) -> bool:
    with _lock:
        teses = _ler()
        novas = [t for t in teses if t.get("id") != tese_id]
        if len(novas) == len(teses):
            return False
        _gravar(novas)
    return True


# ---- importação dos blocos do teses-dpu ------------------------------------------

def _md_para_html(texto: str) -> str:
    """Bloco Markdown simples (citação `>`, **negrito**, *itálico*) → `<p>`."""
    linhas = [re.sub(r"^\s*>\s?", "", l) for l in texto.splitlines()]
    paragrafos, atual = [], []
    for l in linhas:
        if l.strip():
            atual.append(l.strip())
        elif atual:
            paragrafos.append(" ".join(atual))
            atual = []
    if atual:
        paragrafos.append(" ".join(atual))

    def inline(s: str) -> str:
        s = s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        s = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", s)
        s = re.sub(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)", r"<em>\1</em>", s)
        return s
    return "".join(f"<p>{inline(p)}</p>" for p in paragrafos if p)


def importar_teses_dpu() -> dict:
    """Importa os "blocos reutilizáveis" de `teses-dpu/areas/*/modelos/
    00-blocos-reutilizaveis.md` como teses INATIVAS (para o Defensor revisar,
    preencher "quando se aplica" e ativar). Não duplica: um bloco já importado
    (mesma origem) é pulado."""
    teses_existentes = _ler()
    ja = {t.get("origem") for t in teses_existentes}
    novas = 0
    for arq in sorted(TESES_DPU_DIR.glob("areas/*/modelos/00-blocos-reutilizaveis.md")):
        area = arq.parent.parent.name
        area = area if area in _AREAS_VALIDAS else "outra"
        texto = arq.read_text(encoding="utf-8", errors="replace")
        for m in re.finditer(r"^##\s+(Bloco[^\n]*)\n(.*?)(?=^##\s|\Z)", texto, re.M | re.S):
            cabecalho, corpo = m.group(1).strip(), m.group(2)
            corpo = re.sub(r"^\s*---\s*$", "", corpo, flags=re.M).strip()
            titulo = re.sub(r"^Bloco\s+[A-Z0-9]+\s*[—–-]\s*", "", cabecalho).strip() or cabecalho
            origem = f"teses-dpu:{area}:{cabecalho}"
            if origem in ja or not corpo:
                continue
            html = _md_para_html(corpo)
            if not html:
                continue
            salvar({"titulo": titulo, "areas": [area], "quando_aplica": "", "texto_html": html,
                    "fundamentos": "", "ativa": False, "origem": origem})
            ja.add(origem)
            novas += 1
    return {"importadas": novas}


# ---- uso na geração de petições -----------------------------------------------

def _texto_simples(html: str) -> str:
    t = re.sub(r"</p\s*>|<br\s*/?>", "\n", html or "", flags=re.I)
    t = re.sub(r"<[^>]+>", "", t)
    for a, b in (("&nbsp;", " "), ("&amp;", "&"), ("&lt;", "<"), ("&gt;", ">"), ("&quot;", '"')):
        t = t.replace(a, b)
    return re.sub(r"\n{3,}", "\n\n", t).strip()


def bloco_prompt() -> str:
    """Trecho do prompt com as teses ATIVAS e a regra de uso. Vazio se não houver
    nenhuma tese ativa.

    Regra (20/10/2026, pedido do Defensor — a 1ª versão pedia à IA o tópico
    "adaptado ao caso" e ela resumia, reescrevia e cortava ementas/artigos): a IA
    só DECIDE quais teses se aplicam e ONDE entram, escrevendo uma marca
    `[[TESE:id]]` sozinha no parágrafo. O texto em si é inserido pelo painel,
    IPSIS LITTERIS do cadastro (`inserir_teses_html` / `inserir_teses_docx`) —
    nada da tese passa pela redação da IA."""
    ativas = [t for t in _ler() if t.get("ativa")]
    if not ativas:
        return ""
    rot = {a["chave"]: a["rotulo"] for a in AREAS}
    itens = []
    for t in ativas:
        itens.append(
            f"### TESE {t['id']} — {t.get('titulo', '')}\n"
            f"Área(s): {', '.join(rot.get(a, a) for a in t.get('areas') or [])}\n"
            f"Quando se aplica: {t.get('quando_aplica') or '(não informado — avalie pelo título e pelo texto)'}\n"
            + (f"Fundamentos: {t['fundamentos']}\n" if t.get("fundamentos") else "")
            + f"Texto cadastrado (SÓ PARA CONSULTA — o painel o insere sozinho):\n"
              f"{_texto_simples(t.get('texto_html', ''))}\n"
        )
    return (
        "\n\n## BANCO DE TESES DO DEFENSOR (use SÓ as que se aplicarem ao caso)\n"
        "Abaixo estão tópicos cadastrados pelo Defensor — fonte segura, que entra na peça "
        "EXATAMENTE como está. Para CADA tese, confira se o \"Quando se aplica\" corresponde "
        "aos fatos e documentos DESTE caso. Se não corresponder, ignore a tese; não force.\n"
        "COMO INSERIR uma tese que se aplica: escreva, no ponto da peça em que o tópico deve "
        "entrar, UM PARÁGRAFO CONTENDO APENAS a marca `[[TESE:<id>]]` (ex.: "
        "`<p>[[TESE:t1a2b3c4]]</p>`; no .docx, um parágrafo cujo texto é só `[[TESE:t1a2b3c4]]`). "
        "O painel troca a marca pelo texto INTEGRAL cadastrado — título do tópico, parágrafos, "
        "citações, ementas e artigos de lei, sem nenhuma alteração.\n"
        "PROIBIDO: copiar, resumir, parafrasear, reescrever ou complementar o texto da tese; "
        "escrever um título próprio para ela (o texto cadastrado já traz o título); e repetir "
        "os argumentos dela em outro ponto da peça. Você PODE escrever, com suas palavras, uma "
        "frase curta ANTES ou DEPOIS da marca ligando a tese aos fatos do caso (ex.: \"No caso, "
        "a autora recolheu…\"), e fazer nos pedidos a referência ao que a tese sustenta.\n"
        "Não use a marca para texto que não está no banco. NÃO liste, nem comente, quais teses "
        "aplicou ou descartou.\n\n"
        + "\n".join(itens)
    )


# ---- inserção do texto cadastrado no lugar das marcas ------------------------------

_RE_ID = r"([A-Za-z0-9_-]+)"
# Parágrafo que contém só a marca (com eventuais tags de formatação em volta).
_RE_P_ABRE = re.compile(r"<p\b[^>]*>(?:\s|<[^>]+>)*\[\[TESE:" + _RE_ID + r"\]\](?:\s|<[^>]+>)*</p>", re.I)
_RE_ABRE = re.compile(r"\[\[TESE:" + _RE_ID + r"\]\]", re.I)
_RE_P_FECHA = re.compile(r"<p\b[^>]*>(?:\s|<[^>]+>)*\[\[/TESE\]\](?:\s|<[^>]+>)*</p>|\[\[/TESE\]\]", re.I)


def _html_da_tese(tese_id: str) -> str:
    """Texto cadastrado da tese, em azul. "" se o id não existir (marca inventada)."""
    t = obter(tese_id) if tese_id else None
    return _azul(t.get("texto_html") or "") if t else ""


def inserir_teses_html(html: str) -> str:
    """Troca cada marca `[[TESE:id]]` pelo texto INTEGRAL cadastrado da tese, em
    azul. Se a IA ainda assim escrever algo entre `[[TESE:id]]` e `[[/TESE]]`
    (formato antigo), esse trecho é DESCARTADO — vale o cadastro. Marca de id
    desconhecido ou fechamento solto: só somem."""
    if not html or "[[" not in html:
        return html
    saida, pos = [], 0
    while True:
        m = _RE_P_ABRE.search(html, pos) or _RE_ABRE.search(html, pos)
        m2 = _RE_ABRE.search(html, pos)
        if m2 and m and m2.start() < m.start():  # marca solta no meio do texto vem antes
            m = m2
        if not m:
            saida.append(html[pos:])
            break
        antes = html[pos:m.start()]
        inline = not m.group(0).lower().startswith("<p")
        saida.append(antes)
        fim = m.end()
        f = _RE_P_FECHA.search(html, fim)
        prox = _RE_ABRE.search(html, fim)
        if f and (not prox or f.start() < prox.start()):
            fim = f.end()  # descarta o que a IA escreveu entre as marcas
        tese = _html_da_tese(m.group(1))
        # Marca no meio de um parágrafo: fecha o parágrafo, entra a tese, reabre.
        saida.append(f"</p>{tese}<p>" if inline and tese else tese)
        pos = fim
    return _RE_P_FECHA.sub("", "".join(saida))


# Compatibilidade com o nome antigo (ainda chamado por minuta/intimação).
aplicar_cor = inserir_teses_html


def _sem_cor_propria(html: str) -> str:
    """Tira as cores embutidas no cadastro (ex.: cinza/branco herdados do Word na
    importação) — senão sobrepõem o azul e o trecho fica cinza ou invisível. Só
    a cor sai; o texto não muda."""
    return re.sub(r"(?<![-\w])color\s*:\s*[^;\"']+;?\s*", "", html or "", flags=re.I)


def _azul(trecho: str) -> str:
    trecho = _sem_cor_propria(trecho)

    def pinta(m):
        return f'{m.group(1)}<span style="color: {COR_TESE_HEX};">{m.group(2)}</span>{m.group(3)}'
    novo = re.sub(r"(<(?:p|li|blockquote)\b[^>]*>)(.*?)(</(?:p|li|blockquote)>)", pinta, trecho, flags=re.S | re.I)
    if novo == trecho and re.sub(r"<[^>]+>", "", trecho).strip():
        novo = f'<span style="color: {COR_TESE_HEX};">{trecho}</span>'
    return novo


def eh_cor_tese(estilo: str) -> bool:
    """True se o `style` de um elemento traz a cor do banco de teses (hex ou rgb)
    — usado pelo conversor HTML→.docx, que só mantém esta cor."""
    s = (estilo or "").lower().replace(" ", "")
    return "color:#0066cc" in s or "color:rgb(0,102,204)" in s


def _paragrafos_da_tese(doc, tese_html: str) -> list:
    """Cria no fim de `doc` os parágrafos do texto cadastrado da tese, no padrão
    das petições da DPU (peticao_base: Arial 12, recuo de 1ª linha; citação com
    recuo à esquerda e corpo 10; título em caixa alta sem recuo), tudo em azul.
    Devolve os parágrafos criados (o chamador os move para o lugar da marca)."""
    import sys

    from bs4 import BeautifulSoup, NavigableString
    from docx.shared import Cm, RGBColor

    from services.peticao_ia_service import MODELOS_DIR, _extrair_runs
    sys.path.insert(0, str(MODELOS_DIR))
    from peticao_base import FI, L, par, r as escrever_run

    azul = RGBColor(*COR_TESE_RGB)
    criados, contadores = [], {}
    for bloco in BeautifulSoup(tese_html or "", "html.parser").find_all(["p", "li", "blockquote"]):
        if bloco.find(["p", "li", "blockquote"]):
            continue  # contêiner (ex.: <blockquote><p>…</p></blockquote>): vale o filho
        runs = _extrair_runs(bloco, NavigableString)
        texto = "".join(t for t, *_ in runs).strip()
        if not texto:
            continue
        classes = bloco.get("class") or []
        eh_li = bloco.name == "li"
        eh_citacao = (not eh_li) and (
            bloco.name == "blockquote" or bloco.find_parent("blockquote") is not None
            or any(str(c).startswith("ql-indent-") for c in classes))
        eh_titulo = (not eh_li and not eh_citacao and len(texto) <= 220
                     and texto.upper() == texto and any(c.isalpha() for c in texto))
        marcador = ""
        if eh_li:
            pai = bloco.parent
            if pai is not None and getattr(pai, "name", "") == "ol":
                contadores[id(pai)] = contadores.get(id(pai), 0) + 1
                marcador = f"{contadores[id(pai)]}. "
            else:
                marcador = "• "
        if eh_citacao:
            p = par(doc, fi=Cm(0), li=FI + Cm(1))
        elif eh_li:
            p = par(doc, fi=Cm(0), li=Cm(1))
        elif eh_titulo:
            p = par(doc, align=L, fi=Cm(0))
        else:
            p = par(doc, fi=FI)
        for t, b, i, u, _a in ([(marcador, False, False, False, True)] if marcador else []) + runs:
            if t:
                escrever_run(p, t, bold=b, italic=i, underline=u, small=eh_citacao, color=azul)
        criados.append(p)
    return criados


def inserir_teses_docx(docx_path) -> int:
    """Para fluxos em que a própria IA gera o .docx (Ações no e-proc › "Gerar
    petição com IA"; petição inicial do "Analisar PAJ"): cada parágrafo-marca
    `[[TESE:id]]` vira o texto INTEGRAL cadastrado da tese, em azul. Se houver
    texto da IA entre `[[TESE:id]]` e `[[/TESE]]` (formato antigo), ele é
    descartado. Marca de id desconhecido/fechamento solto: só somem. Devolve
    quantas teses foram inseridas."""
    import docx as _docx

    d = _docx.Document(str(docx_path))
    paragrafos = list(d.paragraphs)
    inseridas, remover, i = 0, [], 0
    while i < len(paragrafos):
        p = paragrafos[i]
        m = _RE_ABRE.search(p.text or "")
        if not m:
            if "[[/TESE]]" in (p.text or ""):
                if (p.text or "").strip() == "[[/TESE]]":
                    remover.append(p)
                else:
                    for run in p.runs:
                        run.text = (run.text or "").replace("[[/TESE]]", "")
            i += 1
            continue
        # Texto da IA entre a marca e um [[/TESE]] (antes da próxima marca): descarta.
        j = i + 1
        while j < len(paragrafos):
            tj = paragrafos[j].text or ""
            if _RE_ABRE.search(tj):
                break
            if "[[/TESE]]" in tj:
                remover.extend(paragrafos[i + 1:j + 1])
                i = j
                break
            j += 1
        tese = obter(m.group(1))
        if tese:
            for novo in _paragrafos_da_tese(d, tese.get("texto_html") or ""):
                p._element.addprevious(novo._element)
            inseridas += 1
        sobra = _RE_ABRE.sub("", p.text or "").strip()
        if sobra:  # marca no meio de um parágrafo com outro texto: só tira a marca
            for run in p.runs:
                run.text = _RE_ABRE.sub("", run.text or "")
        else:
            remover.append(p)
        i += 1
    for p in remover:
        with contextlib.suppress(Exception):
            p._element.getparent().remove(p._element)
    d.save(str(docx_path))
    return inseridas


# Compatibilidade com o nome antigo.
colorir_docx = inserir_teses_docx


def destacar_em_pecas(pasta, desde_ts: float) -> list[str]:
    """Para a petição inicial do "Analisar PAJ" (a IA gera o .docx pelo
    pipeline do teses-dpu, numa sessão própria): nos .docx da pasta alterados a
    partir de `desde_ts` que tenham marcas `[[TESE…]]`, insere o texto
    cadastrado das teses (`inserir_teses_docx`) e refaz o PDF de mesmo nome.
    Devolve os arquivos tratados."""
    import docx as _docx
    from pathlib import Path as _P

    tratados = []
    pasta = _P(pasta)
    if not pasta.exists():
        return tratados
    for arq in pasta.glob("*.docx"):
        with contextlib.suppress(Exception):
            if arq.stat().st_mtime < desde_ts or arq.name.startswith("~$"):
                continue
            if not any("[[TESE" in (p.text or "") or "[[/TESE" in (p.text or "")
                       for p in _docx.Document(str(arq)).paragraphs):
                continue
            inserir_teses_docx(arq)
            from services.peticao_ia_service import _docx_para_pdf
            _docx_para_pdf(arq, arq.with_suffix(".pdf"))
            tratados.append(arq.name)
    return tratados


def run_eh_tese(run) -> bool:
    """True se o run do .docx está no azul do banco de teses."""
    with contextlib.suppress(Exception):
        cor = run.font.color
        return cor is not None and cor.type is not None and str(cor.rgb).upper() == "0066CC"
    return False


def excluir_varios(ids: list[str]) -> int:
    """Exclui as teses de `ids` de uma vez (botão "Excluir as listadas" da página).
    Devolve quantas saíram."""
    alvo = set(ids or [])
    if not alvo:
        return 0
    with _lock:
        teses = _ler()
        novas = [t for t in teses if t.get("id") not in alvo]
        _gravar(novas)
    return len(teses) - len(novas)


# ---- importação de arquivos (PDF / DOCX) -----------------------------------------
# Botão "Nova tese" › "Importar arquivo" (1 arquivo = 1 tese, abre o formulário
# já preenchido para revisar) ou "Importar pasta" (cada arquivo vira uma tese
# INATIVA, para revisar, preencher "Quando se aplica" e ativar) — 19/10/2026.

def _esc(t: str) -> str:
    return (t or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _titulo_do_nome(nome: str) -> str:
    base = re.sub(r"\.(docx|pdf)$", "", Path(nome).name, flags=re.I)
    return re.sub(r"[_\s]+", " ", base).strip()[:200]


def _parece_titulo(texto: str, negrito: bool = False, estilo: str = "") -> bool:
    t = (texto or "").strip()
    if not t or len(t) > 160 or t.endswith((".", ";", ",")) and len(t) > 60:
        return False
    return (estilo.lower().startswith(("heading", "título", "titulo", "title"))
            or negrito or (t.upper() == t and any(c.isalpha() for c in t)))


def _ler_docx(conteudo: bytes) -> tuple[str, str]:
    import io

    import docx as _docx
    from docx.shared import Cm

    d = _docx.Document(io.BytesIO(conteudo))
    titulo, partes = "", []
    for p in d.paragraphs:
        texto = (p.text or "").strip()
        if not texto:
            continue
        runs = [r for r in p.runs if (r.text or "").strip()]
        todo_negrito = bool(runs) and all(r.bold for r in runs)
        if not titulo and not partes and _parece_titulo(texto, todo_negrito, getattr(p.style, "name", "") or ""):
            titulo = texto[:200]
            continue
        html = ""
        for r in p.runs:
            t = _esc(r.text)
            if not t:
                continue
            if r.bold:
                t = f"<strong>{t}</strong>"
            if r.italic:
                t = f"<em>{t}</em>"
            if r.underline:
                t = f"<u>{t}</u>"
            html += t
        recuo = p.paragraph_format.left_indent
        tag = "blockquote" if (recuo is not None and recuo >= Cm(2)) else "p"
        partes.append(f"<{tag}>{html.strip()}</{tag}>")
    return titulo, "".join(partes)


def _ler_pdf(conteudo: bytes) -> tuple[str, str]:
    import fitz

    blocos = []
    with fitz.open(stream=conteudo, filetype="pdf") as doc:
        for pagina in doc:
            for b in pagina.get_text("blocks"):
                texto = re.sub(r"\s*\n\s*", " ", b[4] or "").strip()
                if texto and not re.fullmatch(r"\d{1,4}", texto):  # pula nº de página solto
                    blocos.append(texto)
    if not blocos:
        raise ValueError("PDF sem texto legível (digitalizado?) — use um PDF com texto ou o .docx.")
    titulo = ""
    if _parece_titulo(blocos[0]):
        titulo = blocos.pop(0)[:200]
    return titulo, "".join(f"<p>{_esc(t)}</p>" for t in blocos)


def ler_arquivo(nome: str, conteudo: bytes) -> dict:
    """{titulo, texto_html} lidos de um .docx ou .pdf (sem gravar nada). Título:
    o 1º parágrafo se parecer título (estilo de título, negrito ou maiúsculas);
    senão, o nome do arquivo."""
    ext = Path(nome or "").suffix.lower()
    if ext == ".docx":
        titulo, html = _ler_docx(conteudo)
    elif ext == ".pdf":
        titulo, html = _ler_pdf(conteudo)
    else:
        raise ValueError("Formato não aceito — use .docx ou .pdf.")
    if not re.sub(r"<[^>]+>", "", html).strip():
        raise ValueError("O arquivo não tem texto para o tópico.")
    return {"titulo": titulo or _titulo_do_nome(nome), "texto_html": html}


def importar_arquivos(arquivos: list[tuple[str, bytes]]) -> dict:
    """Importação em bloco: cada .docx/.pdf vira uma tese INATIVA (área "geral").
    Arquivo já importado antes (mesmo nome) é pulado. Outros formatos são
    ignorados com o motivo."""
    ja = {t.get("origem") for t in _ler()}
    importadas, ignoradas = 0, []
    for nome, conteudo in arquivos:
        base = Path(nome or "").name
        if base.startswith("~$"):
            continue  # arquivo temporário do Word
        origem = f"arquivo:{base}"
        if origem in ja:
            ignoradas.append({"arquivo": base, "motivo": "já importado antes"})
            continue
        try:
            lido = ler_arquivo(base, conteudo)
        except ValueError as e:
            ignoradas.append({"arquivo": base, "motivo": str(e)})
            continue
        except Exception as e:  # noqa: BLE001
            ignoradas.append({"arquivo": base, "motivo": f"não foi possível ler ({type(e).__name__})"})
            continue
        salvar({**lido, "areas": ["geral"], "quando_aplica": "", "fundamentos": "",
                "ativa": False, "origem": origem})
        ja.add(origem)
        importadas += 1
    return {"importadas": importadas, "ignoradas": ignoradas}
