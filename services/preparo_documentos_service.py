"""Preparação dos documentos do PAJ ANTES de qualquer análise com IA.

Decisão do Defensor (13/10/2026): ao determinar uma análise, o painel confere
que TODOS os documentos do PAJ estão na pasta; o que ainda faltar é baixado
automaticamente (MCP do SIS-DPU 2.0); com todos conferidos, passa OCR em tudo o
que estiver sem texto e só então prossegue. Sem modal: abrir o PAJ já baixa os
documentos (static/js/ao_vivo.js), então normalmente só sobra o OCR.

Cascata de texto por arquivo (gera o `.txt` ao lado do arquivo, o mesmo
"companion" que o PROMPT_MAX e as rotinas de IA já leem):
  - PDF: texto próprio do arquivo; se não houver, OCR local (Tesseract) por
    página — `ingestao.ocr.extrair_texto`;
  - imagem (jpg/jpeg/png/gif): convertida em PDF de 1 página e idem;
  - Word (docx): texto do documento (python-docx).
O que continuar fraco segue para `ocr_llm_service.melhorar_ocr_paj` (ocrmypdf
e, por último, transcrição por IA), chamado por quem dispara a análise.
"""

from __future__ import annotations

import asyncio
import concurrent.futures
import json
import logging
import re
import tempfile
import time
from pathlib import Path

from config import PAJS_DIR

log = logging.getLogger("preparo_documentos")

_EXT_IMAGEM = {".jpg", ".jpeg", ".png", ".gif"}
_EXT_TEXTO = {".pdf", ".docx"} | _EXT_IMAGEM


def _rodar(coro):
    """Roda uma corrotina de qualquer contexto (thread de análise ou não)."""
    with concurrent.futures.ThreadPoolExecutor(max_workers=1) as ex:
        return ex.submit(asyncio.run, coro).result()


def _documentos_listados(paj_norm: str) -> list[dict]:
    """Todos os documentos das movimentações (metadata), no formato de download."""
    pasta = PAJS_DIR / paj_norm
    try:
        m = json.loads((pasta / "metadata.json").read_text(encoding="utf-8"))
    except Exception:  # noqa: BLE001
        return []
    itens = []
    for mv in (m.get("detalhes_sisdpu") or {}).get("movimentacoes") or []:
        for d in mv.get("documentos") or []:
            href = d.get("urlDownload") or ""
            if re.search(r"/(\d+)/[^/]*$", href):
                itens.append({"seq": mv.get("seq"), "data": mv.get("data"),
                              "nomeArq": d.get("nomeArquivo") or "arquivo", "href": href})
    return itens


def _ids_locais(pasta_arquivos: Path) -> set[str]:
    if not pasta_arquivos.is_dir():
        return set()
    return {m.group(1) for f in pasta_arquivos.iterdir()
            if f.is_file() and f.suffix.lower() != ".txt"
            for m in [re.match(r"seq\d+ - (\d+) - ", f.name)] if m}


def _aguardar_download_em_andamento(paj_norm: str, log_fn, limite_seg: int = 900) -> None:
    """Se o download ao abrir o PAJ ainda está rodando, espera terminar (evita
    baixar o mesmo arquivo duas vezes ao mesmo tempo)."""
    try:
        from services import ao_vivo_service
    except Exception:  # noqa: BLE001
        return
    inicio = time.time()
    avisou = False
    while time.time() - inicio < limite_seg:
        p = ao_vivo_service.progresso_anexos(paj_norm) or {}
        if not p.get("rodando"):
            return
        if not avisou:
            log_fn("  [preparo] aguardando o download dos documentos iniciado ao abrir o PAJ…")
            avisou = True
        time.sleep(3)


def _texto_de(arquivo: Path) -> str:
    from ingestao import ocr

    ext = arquivo.suffix.lower()
    if ext == ".pdf":
        return ocr.extrair_texto(arquivo)
    if ext == ".docx":
        import docx  # python-docx

        d = docx.Document(str(arquivo))
        partes = [p.text for p in d.paragraphs]
        for t in d.tables:
            for linha in t.rows:
                partes.append(" | ".join(c.text for c in linha.cells))
        return "\n".join(x for x in partes if x.strip())
    if ext in _EXT_IMAGEM:
        import fitz  # PyMuPDF

        with fitz.open(str(arquivo)) as img:
            pdf_bytes = img.convert_to_pdf()
        with tempfile.TemporaryDirectory() as tmp:
            tmp_pdf = Path(tmp) / "imagem.pdf"
            tmp_pdf.write_bytes(pdf_bytes)
            return ocr.extrair_texto(tmp_pdf)
    return ""


def preparar_documentos_paj(paj_norm: str, log_fn=None) -> dict:
    """Confere/baixa/OCR. Best-effort: nunca impede a análise por erro próprio —
    registra as pendências no resultado e no log."""

    def _log(msg: str) -> None:
        log.info(msg)
        if log_fn:
            try:
                log_fn(msg)
            except Exception:  # noqa: BLE001
                pass

    res = {"listados": 0, "baixados": 0, "ocr_feitos": [], "pendencias": []}
    pasta = PAJS_DIR / paj_norm
    arquivos = pasta / "Arquivos"
    if not pasta.is_dir():
        return res

    # 1) conferência + download do que faltar (fonte SIS-DPU 2.0)
    listados = _documentos_listados(paj_norm)
    res["listados"] = len(listados)
    try:
        from ingestao import sisdpu2_mcp_client as sis2
        from services import config_painel_service
        mcp = config_painel_service.get_fonte_dados() == "mcp" and sis2.conectado()
    except Exception:  # noqa: BLE001
        mcp = False
    if mcp and listados:
        _aguardar_download_em_andamento(paj_norm, _log)
        locais = _ids_locais(arquivos)
        faltam = [d for d in listados if re.search(r"/(\d+)/[^/]*$", d["href"]).group(1) not in locais]
        if faltam:
            _log(f"  [preparo] {len(faltam)} documento(s) ainda não baixado(s) — baixando pelo SIS-DPU 2.0…")
            from ingestao import sisdpu2_fonte as fonte
            try:
                r = _rodar(fonte.baixar_todos_arquivos(arquivos, faltam, log=_log))
                res["baixados"] = r.get("baixados", 0)
                for e in r.get("erros") or []:
                    res["pendencias"].append(f"download: {e}")
            except Exception as e:  # noqa: BLE001
                res["pendencias"].append(f"download: {type(e).__name__}: {e}")

    # 2) OCR/texto de tudo o que estiver sem `.txt`
    if arquivos.is_dir():
        for f in sorted(arquivos.iterdir()):
            if not f.is_file() or f.suffix.lower() not in _EXT_TEXTO:
                continue
            companion = f.with_suffix(".txt")
            if companion.exists() and companion.stat().st_size > 0:
                continue
            try:
                texto = _texto_de(f)
            except Exception as e:  # noqa: BLE001
                res["pendencias"].append(f"OCR: {f.name}: {type(e).__name__}: {e}")
                continue
            if texto and texto.strip():
                companion.write_text(texto, encoding="utf-8")
                res["ocr_feitos"].append(f.name)
            else:
                res["pendencias"].append(f"OCR: {f.name}: sem texto extraído")
    _log(f"  [preparo] {res['listados']} documento(s) no PAJ; {res['baixados']} baixado(s) agora; "
         f"texto extraído de {len(res['ocr_feitos'])}; pendências: {len(res['pendencias'])}")
    for p in res["pendencias"][:20]:
        _log(f"    - {p}")
    return res
