"""Runner do cálculo do valor da causa — executado como SUBPROCESS pelo
`valor_causa_service`, sob o Python do sistema (que tem python-docx + pywin32).

Isola o Word/COM e a dependência de docx do processo do painel (cujo venv não
tem essas libs). Faz duas coisas: gera o PDF da tabela de prestações (via
`calc_valor_causa.gerar_pdf`) e injeta o total no texto da petição (placeholder
no `.txt` e no `.docx`).

Uso:
    python valor_causa_runner.py <job.json>

job.json:
    {
      "calc_dir":  "C:\\Claude\\DPU\\Modelos\\Peticoes",
      "dados":     { ...entrada do calc_valor_causa... },
      "pdf_path":  "...\\CALC VALOR DA CAUSA.pdf",
      "placeholder": "{{VALOR_CAUSA}}",
      "txt_files":  ["...\\peticao_inicial.txt"],
      "docx_files": ["...\\peticao_inicial.docx"]
    }

Saída (stdout, ÚLTIMA linha, prefixada por RESULT_JSON:):
    RESULT_JSON:{"ok": true, "total_nominal": ..., "total_atualizado": ...,
                 "valor_causa": "R$ ...", "injetado_em": [...]}
"""

import contextlib
import json
import sys
from pathlib import Path


# ── Valor por extenso (pt-BR, moeda) ──────────────────────────────────────
# Padrão de redação deste ofício: o campo do valor da causa sai SEMPRE como o
# valor numérico seguido do valor por extenso entre parênteses. Como é o painel
# que injeta o valor no lugar do marcador {{VALOR_CAUSA}}, o por extenso é gerado
# aqui, a partir do total calculado (nunca pelo LLM). num2words não é dependência
# do projeto, por isso a conversão é local.
_UNI = ['zero', 'um', 'dois', 'três', 'quatro', 'cinco', 'seis', 'sete', 'oito',
        'nove', 'dez', 'onze', 'doze', 'treze', 'catorze', 'quinze', 'dezesseis',
        'dezessete', 'dezoito', 'dezenove']
_DEZ = ['', '', 'vinte', 'trinta', 'quarenta', 'cinquenta', 'sessenta',
        'setenta', 'oitenta', 'noventa']
_CEM = ['', 'cento', 'duzentos', 'trezentos', 'quatrocentos', 'quinhentos',
        'seiscentos', 'setecentos', 'oitocentos', 'novecentos']
_ESCALA = [('', ''), ('mil', 'mil'), ('milhão', 'milhões'),
           ('bilhão', 'bilhões'), ('trilhão', 'trilhões')]


def _grupo_extenso(n: int) -> str:
    """Número de 0 a 999 por extenso ('' para zero)."""
    if n == 0:
        return ''
    if n == 100:
        return 'cem'
    centena, resto = divmod(n, 100)
    partes = []
    if centena:
        partes.append(_CEM[centena])
    if resto:
        if resto < 20:
            partes.append(_UNI[resto])
        else:
            dez, uni = divmod(resto, 10)
            partes.append(_DEZ[dez] + (' e ' + _UNI[uni] if uni else ''))
    return ' e '.join(partes)


def _inteiro_extenso(n: int) -> str:
    """Inteiro não negativo por extenso (até trilhões)."""
    if n == 0:
        return 'zero'
    grupos = []
    escala = 0
    resto = n
    while resto > 0:
        grupos.append((escala, resto % 1000))
        resto //= 1000
        escala += 1
    grupos = [(e, v) for (e, v) in grupos if v != 0]
    grupos.sort(reverse=True)
    textos = []
    for e, v in grupos:
        if e == 0:
            t = _grupo_extenso(v)
        elif e == 1:
            t = 'mil' if v == 1 else _grupo_extenso(v) + ' mil'
        else:
            sing, plur = _ESCALA[e]
            t = _grupo_extenso(v) + ' ' + (sing if v == 1 else plur)
        textos.append((v, t))
    resultado = textos[0][1]
    for i in range(1, len(textos)):
        v, t = textos[i]
        ultimo = i == len(textos) - 1
        conector = ' e ' if (ultimo and (v < 100 or v % 100 == 0)) else ', '
        resultado += conector + t
    return resultado


def _extenso_reais(valor: float) -> str:
    """Valor monetário por extenso, ex.: 34181.90 -> 'trinta e quatro mil,
    cento e oitenta e um reais e noventa centavos'."""
    valor = round(float(valor) + 1e-9, 2)
    reais = int(valor)
    centavos = int(round((valor - reais) * 100))
    if centavos >= 100:
        reais += 1
        centavos -= 100
    partes = []
    if reais > 0:
        partes.append(_inteiro_extenso(reais) + (' real' if reais == 1 else ' reais'))
    if centavos > 0:
        cent = _inteiro_extenso(centavos) + (' centavo' if centavos == 1 else ' centavos')
        partes.append(('e ' + cent) if reais > 0 else cent)
    return ' '.join(partes) if partes else 'zero reais'


def _injetar_txt(caminho: Path, placeholder: str, valor: str) -> bool:
    with contextlib.suppress(Exception):
        c = caminho.read_text(encoding="utf-8", errors="replace")
        if placeholder in c:
            caminho.write_text(c.replace(placeholder, valor), encoding="utf-8")
            return True
    return False


def _injetar_docx(caminho: Path, placeholder: str, valor: str) -> bool:
    """Caminho rápido: placeholder num único run (preserva a formatação).
    Fallback: placeholder quebrado entre runs -> concatena no primeiro run."""
    try:
        from docx import Document
    except Exception:
        return False
    with contextlib.suppress(Exception):
        doc = Document(str(caminho))
        n = 0

        def _proc(par) -> None:
            nonlocal n
            for r in par.runs:
                if placeholder in r.text:
                    r.text = r.text.replace(placeholder, valor)
                    n += 1
                    return
            if placeholder in par.text and par.runs:
                novo = par.text.replace(placeholder, valor)
                for r in par.runs:
                    r.text = ""
                par.runs[0].text = novo
                n += 1

        for p in doc.paragraphs:
            _proc(p)
        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        _proc(p)
        if n:
            doc.save(str(caminho))
            return True
    return False


def main() -> int:
    if len(sys.argv) < 2:
        print('RESULT_JSON:{"ok": false, "erro": "job.json não informado"}')
        return 1
    job = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))

    # sys.path para importar calc_valor_causa + peticao_base (pasta-mãe)
    calc_dir = job["calc_dir"]
    sys.path.insert(0, calc_dir)
    sys.path.insert(0, str(Path(calc_dir).parent))

    try:
        import calc_valor_causa as calc
    except Exception as e:
        print('RESULT_JSON:' + json.dumps({"ok": False, "erro": f"import calc: {e}"}))
        return 1

    dados = job["dados"]
    pdf_path = job["pdf_path"]

    # calc.gerar_pdf imprime no stdout — redireciona pro stderr pra não poluir o
    # JSON de saída.
    try:
        with contextlib.redirect_stdout(sys.stderr):
            total_nom, total_atual = calc.gerar_pdf(dados, pdf_path)
    except Exception as e:
        print('RESULT_JSON:' + json.dumps({"ok": False, "erro": f"gerar_pdf: {e}"}))
        return 1

    valor_final = total_atual or total_nom
    # Padrão deste ofício: valor numérico seguido do por extenso entre parênteses.
    valor_str = calc._moeda(valor_final) + f" ({_extenso_reais(valor_final)})"

    placeholder = job.get("placeholder", "{{VALOR_CAUSA}}")
    injetados = []
    for t in job.get("txt_files", []):
        if _injetar_txt(Path(t), placeholder, valor_str):
            injetados.append(Path(t).name)
    for d in job.get("docx_files", []):
        if _injetar_docx(Path(d), placeholder, valor_str):
            injetados.append(Path(d).name)

    print("RESULT_JSON:" + json.dumps({
        "ok": True,
        "total_nominal": total_nom,
        "total_atualizado": total_atual,
        "valor_causa": valor_str,
        "injetado_em": injetados,
    }, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
